from __future__ import annotations

import io
import zipfile
from datetime import datetime
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from visitacampo.forms.answers import validate_answers
from visitacampo.forms.definition import RepeatField
from visitacampo.forms.visibility import Comparison
from visitacampo.forms.xlsform import InvalidXlsForm, import_xlsform


def workbook(survey, choices=None, settings=None, extra=None):
    book = Workbook()
    book.remove(book.active)
    for name, rows in {
        "survey": survey,
        "choices": choices,
        "settings": settings,
        **(extra or {}),
    }.items():
        if rows is not None:
            sheet = book.create_sheet(name)
            for row in rows:
                sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def test_roundtrip_defaults_repeats_choices_qr_and_visibility():
    result = import_xlsform(
        workbook(
            [
                ["type", "name", "label", "default", "required", "relevant", "hint"],
                ["text", "site", "Sitio", "Bogotá", "yes", None, "Nombre del sitio"],
                ["begin repeat", "assets", "Activos"],
                ["barcode", "code", "Código", None, "yes"],
                ["select_one state", "state", "Estado", "ok"],
                ["integer", "count", "Cantidad", 3, None, "${state} = 'ok'"],
                ["select_multiple state", "many", "Estados", "ok bad"],
                ["end repeat"],
            ],
            [["list_name", "name", "label"], ["state", "ok", "Bueno"], ["state", "bad", "Malo"]],
            [["form_title", "form_id", "version"], ["Inspección", "inspection", "1"]],
        )
    )
    assert result.title == {"es": "Inspección"}
    assert result.fields[0].default == "Bogotá"
    assert isinstance(result.fields[1], RepeatField)
    assert result.fields[1].fields[2].default == 3
    assert result.fields[1].fields[3].default == ["ok", "bad"]
    assert (
        validate_answers(
            result, {"site": "Bogotá", "assets": [{"code": "A-1", "state": "ok", "count": 3}]}
        )
        == []
    )
    violations = validate_answers(result, {"site": "Bogotá", "assets": [{"state": "bad"}]})
    assert [(v.field, v.reason) for v in violations] == [("answers.assets.0.code", "required")]


@pytest.mark.parametrize(
    "kind,expected",
    [
        ("text", "text"),
        ("integer", "integer"),
        ("decimal", "number"),
        ("date", "date"),
        ("dateTime", "datetime"),
        ("geopoint", "location"),
        ("geotrace", "line"),
        ("geoshape", "polygon"),
        ("image", "photo"),
        ("barcode", "qr"),
    ],
)
def test_supported_scalar_types(kind, expected):
    result = import_xlsform(workbook([["type", "name", "label"], [kind, "value", "Valor"]]))
    assert result.fields[0].type == expected


def test_signature_appearance_and_numeric_literal():
    parsed = import_xlsform(
        workbook(
            [
                ["type", "name", "label", "appearance", "default", "relevant"],
                ["image", "sign", "Firma", "signature"],
                ["decimal", "n", "Número", None, 2.5],
                ["text", "note", "Nota", None, None, "${n} >= 2.5"],
            ]
        )
    )
    assert parsed.fields[0].type == "signature"
    assert parsed.fields[1].default == 2.5
    assert isinstance(parsed.fields[2].visible_if, Comparison)
    assert parsed.fields[2].visible_if.value == 2.5


@pytest.mark.parametrize(
    "rows",
    [
        [["type", "name", "label", "constraint"], ["text", "x", "X", ". > 2"]],
        [["type", "name", "label", "required"], ["text", "x", "X", "${x} = 1"]],
        [["type", "name", "label", "default"], ["text", "x", "X", "today()"]],
        [["type", "name", "label", "default"], ["text", "x", "X", "=1+1"]],
        [["type", "name", "label", "default"], ["integer", "x", "X", 1.5]],
        [["type", "name", "label", "relevant"], ["text", "x", "X", "true()"]],
        [["type", "name", "label", "relevant"], ["text", "x", "X", "${x} = today()"]],
        [["type", "name", "label", "appearance"], ["text", "x", "X", "hidden"]],
        [["type", "name", "label"], ["calculate", "x", "X"]],
        [["type", "name", "label"], ["text other", "x", "X"]],
        [["type", "name", "label"], ["select_one missing", "x", "X"]],
        [["type", "name", "label"], ["end_repeat"]],
        [["type", "name", "label"], ["begin_repeat", "x", "X"]],
        [["type", "name", "label"], ["text", "x", "X"], ["text", "x", "X"]],
        [["type", "name", "name"], ["text", "x", "X"]],
        [["type", "name", "label", None], ["text", "x", "X", "unheaded"]],
        [
            ["type", "name", "label", "relevant"],
            ["text", "a", "A"],
            ["text", "x", "X", "${a} = 'A' and ${a} = 'B'"],
        ],
        [
            ["type", "name", "label", "relevant"],
            ["text", "a", "A"],
            ["text", "x", "X", "${a} = 'A' or ${a} = \"B\""],
        ],
        [["type", "name", "label", "default"], ["text", "x", "X", "concat('a', 'b')"]],
        [["type", "name", "label", "default"], ["date", "x", "X", datetime(2024, 1, 2, 8, 30)]],
        [["type", "name", "label", "default"], ["integer", "x", "X", datetime(2024, 1, 2)]],
        [["type", "name", "label", "default"], ["integer", "x", "X", True]],
        [["type", "name", "label", "repeat_count"], ["text", "x", "X", 3]],
        [
            ["type", "name", "label", "repeat_count"],
            ["begin repeat", "r", "R", "${n}"],
            ["text", "x", "X"],
            ["end repeat"],
        ],
        [
            ["type", "name", "label"],
            ["begin repeat", "outer", "O"],
            ["begin repeat", "inner", "I"],
            ["text", "x", "X"],
            ["end repeat"],
            ["end repeat"],
        ],
    ],
)
def test_rules_are_never_silently_lost(rows):
    with pytest.raises(InvalidXlsForm):
        import_xlsform(workbook(rows))


def test_plain_text_defaults_with_parentheses_are_literals():
    parsed = import_xlsform(
        workbook([["type", "name", "label", "default"], ["text", "x", "X", "Lote 3 (A)"]])
    )
    assert parsed.fields[0].default == "Lote 3 (A)"


def test_excel_date_cell_default_becomes_iso_date():
    parsed = import_xlsform(
        workbook([["type", "name", "label", "default"], ["date", "x", "X", datetime(2024, 1, 2)]])
    )
    assert parsed.fields[0].default == "2024-01-02"


def test_repeat_count_bounds_nested_repeats():
    parsed = import_xlsform(
        workbook(
            [
                ["type", "name", "label", "repeat_count"],
                ["begin repeat", "outer", "O", 10],
                ["begin repeat", "middle", "M", 10],
                ["begin repeat", "inner", "I", 10],
                ["text", "x", "X"],
                ["end repeat"],
                ["end repeat"],
                ["end repeat"],
            ]
        )
    )
    outer = parsed.fields[0]
    assert isinstance(outer, RepeatField)
    middle = outer.fields[0]
    assert isinstance(middle, RepeatField)
    inner = middle.fields[0]
    assert isinstance(inner, RepeatField)
    assert (outer.max_rows, middle.max_rows, inner.max_rows) == (10, 10, 10)


def test_archive_limits_and_bad_files():
    data = workbook([["type", "name", "label"], ["text", "x", "X"]])
    with patch("visitacampo.forms.xlsform.MAX_FILE_BYTES", 10), pytest.raises(InvalidXlsForm):
        import_xlsform(data)
    with patch("visitacampo.forms.xlsform.MAX_EXPANDED_BYTES", 10), pytest.raises(InvalidXlsForm):
        import_xlsform(data)
    with pytest.raises(InvalidXlsForm):
        import_xlsform(b"not a workbook")
    for entry in ("xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"):
        buffer = io.BytesIO(data)
        with zipfile.ZipFile(buffer, "a") as archive:
            archive.writestr(entry, "blocked")
        with pytest.raises(InvalidXlsForm):
            import_xlsform(buffer.getvalue())


@pytest.mark.parametrize(
    "kwargs",
    [
        {"extra": {"calculations": [["x"]]}},
        {"choices": [["list_name", "name"], ["list", "x"]]},
        {"settings": [["form_title"], ["One"], ["Two"]]},
        {"settings": [["submission_url"], ["https://example.test"]]},
    ],
)
def test_bad_sheets(kwargs):
    with pytest.raises(InvalidXlsForm):
        import_xlsform(workbook([["type", "name", "label"], ["text", "x", "X"]], **kwargs))


def test_missing_sheet_and_dimensions():
    with pytest.raises(InvalidXlsForm):
        import_xlsform(workbook(None, settings=[["form_title"], ["X"]]))
    with patch("visitacampo.forms.xlsform.MAX_SHEET_ROWS", 2), pytest.raises(InvalidXlsForm):
        import_xlsform(
            workbook([["type", "name", "label"], ["text", "a", "A"], ["text", "b", "B"]])
        )


@pytest.mark.parametrize(
    "needle,replacement",
    [
        (b'r="2"', b'r="1000000000"'),
        (b'r="A2"', b'r="ZZ2"'),
        (b'r="A2"', b'r="A1000000000"'),
    ],
)
def test_sparse_xml_is_rejected_before_workbook_iteration(needle, replacement):
    data = workbook([["type", "name", "label"], ["text", "x", "X"]])
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            contents = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                contents = contents.replace(needle, replacement)
            target.writestr(item.filename, contents)
    with patch("visitacampo.forms.xlsform.load_workbook") as load, pytest.raises(InvalidXlsForm):
        import_xlsform(buffer.getvalue())
    load.assert_not_called()
    with pytest.raises(InvalidXlsForm):
        import_xlsform(
            workbook(
                [
                    ["type", "name", "label"] + [None] * 29 + ["constraint"],
                    ["text", "a", "A"] + [None] * 29 + [". > 1"],
                ]
            )
        )
